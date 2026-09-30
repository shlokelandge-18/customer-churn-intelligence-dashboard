"""
PowerPoint Deck Generator for Customer Churn Prediction Project.
Creates a professional 16:9 presentation deck with charts and tables.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    base_dir = Path(__file__).parent
    output_pptx = base_dir / "Customer_Churn_Presentation.pptx"
    chart_image = base_dir / "outputs" / "evaluation_dashboard.png"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    PRIMARY = RGBColor(15, 23, 42)       # Slate Dark #0f172a
    ACCENT = RGBColor(14, 165, 233)      # Sky Blue #0ea5e9
    SECONDARY = RGBColor(71, 85, 105)    # Slate Gray #475569
    TEXT_DARK = RGBColor(30, 41, 59)     # Dark Text
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category="CUSTOMER INTELLIGENCE ML"):
        # Category pill/subtitle
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY
    bg1.line.color.rgb = PRIMARY

    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11), Inches(2.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "DATA SCIENCE & MACHINE LEARNING PROJECT"
    p_tag.font.size = Pt(14)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT
    p_tag.space_after = Pt(12)

    p_title = tf1.add_paragraph()
    p_title.text = "Customer Churn Prediction &\nRisk Intelligence System"
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.text = "An end-to-end lightweight pipeline for real-time attrition risk scoring & proactive retention"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)

    # Footer note
    footer_box = slide1.shapes.add_textbox(Inches(1.2), Inches(6.0), Inches(10), Inches(0.8))
    tf_f = footer_box.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "Presenter: Shlok  |  Technology: Python, Scikit-Learn, Pandas, Flask"
    pf.font.size = Pt(13)
    pf.font.color.rgb = RGBColor(203, 213, 225)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Business Value
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem Statement & Business Motivation", "The Business Context")

    # Card 1: The Cost of Churn
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = LIGHT_BG
    c1.line.color.rgb = RGBColor(226, 232, 240)
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. The Attrition Dilemma\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    bullets1 = [
        "Acquiring new customers is 5x to 7x more costly than retaining existing ones.",
        "Revenue loss accumulates exponentially with recurring subscription models.",
        "Customer churn happens silently before formal cancellation."
    ]
    for b in bullets1:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(10)

    # Card 2: Traditional Limitations
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = LIGHT_BG
    c2.line.color.rgb = RGBColor(226, 232, 240)
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Traditional Shortcomings\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    bullets2 = [
        "Reactive outreach occurs after customer requests account termination.",
        "Manual rule heuristics fail to capture non-linear multidimensional indicators.",
        "Resource-heavy ML models are slow, expensive, and unsuited for lightweight deployment."
    ]
    for b in bullets2:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(10)

    # Card 3: Our Solution
    c3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8))
    c3.fill.solid()
    c3.fill.fore_color.rgb = RGBColor(238, 246, 255)
    c3.line.color.rgb = ACCENT
    tf = c3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Proposed AI System\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    bullets3 = [
        "Lightweight Pipeline: Scikit-learn based, sub-second execution on standard CPU.",
        "Real-Time Risk Scoring: Generates continuous probability & categorical risk tiers.",
        "Prescriptive Retention: Suggests personalized retention tactics to save high-risk accounts."
    ]
    for b in bullets3:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: Dataset Architecture & Features
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Dataset Features & Engineering", "Data Overview")

    left_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = LIGHT_BG
    left_box.line.color.rgb = RGBColor(226, 232, 240)
    tf = left_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Dataset Characteristics\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    details = [
        "Total Sample Size: 1,200 Customer Profiles",
        "Target Variable: Churn (1 = Churned, 0 = Retained)",
        "Class Balance: 33.2% Positive Churn Rate (Reflects industry distribution)",
        "Train / Test Split: 80% Training (960 rows) / 20% Test (240 rows)",
        "Stratification: Preserves identical class ratio in train and test splits to prevent evaluation skew."
    ]
    for d in details:
        pb = tf.add_paragraph()
        pb.text = "✔ " + d
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(12)

    right_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = LIGHT_BG
    right_box.line.color.rgb = RGBColor(226, 232, 240)
    tf = right_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Feature Hierarchy\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    feat_groups = [
        ("Demographic", "Age (18 - 70)"),
        ("Tenure & Contract", "TenureMonths (1 - 72 mos), Contract (Month-to-month, 1-yr, 2-yr)"),
        ("Services Active", "InternetService (Fiber optic, DSL, None), TechSupport, OnlineSecurity"),
        ("Billing & Payment", "MonthlyCharges, TotalCharges, PaperlessBilling, PaymentMethod"),
        ("Preprocessing", "StandardScaler for numeric metrics + OneHotEncoder(drop='first') for categories")
    ]
    for cat, desc in feat_groups:
        pb = tf.add_paragraph()
        pb.text = f"• {cat}: {desc}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: Model Benchmarking & Performance
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Machine Learning Model Benchmarking", "Evaluation & Selection")

    # Table
    rows = 4
    cols = 6
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(2.2)

    table_shape = slide4.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    headers = ["Model", "Algorithm Type", "Accuracy", "Precision", "Recall", "ROC-AUC"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.size = Pt(14)
            p.alignment = PP_ALIGN.CENTER

    data_rows = [
        ["Logistic Regression ★", "Linear Probabilistic", "76.7%", "71.4%", "50.0%", "0.797"],
        ["Random Forest", "Ensemble Bagging", "74.2%", "70.5%", "38.8%", "0.761"],
        ["Gradient Boosting", "Ensemble Boosting", "72.9%", "60.9%", "52.5%", "0.760"],
    ]

    for i, row in enumerate(data_rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = RGBColor(238, 246, 255)
            else:
                cell.fill.fore_color.rgb = WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.alignment = PP_ALIGN.CENTER
                if i == 0:
                    p.font.bold = True
                    p.font.color.rgb = ACCENT
                else:
                    p.font.color.rgb = TEXT_DARK

    # Insights card below table
    note_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.4), Inches(11.7), Inches(2.3))
    note_box.fill.solid()
    note_box.fill.fore_color.rgb = LIGHT_BG
    note_box.line.color.rgb = RGBColor(226, 232, 240)
    tf = note_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Why Logistic Regression Emerged as the Champion:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    reasons = [
        "Superior Discrimination: Highest ROC-AUC (0.797) demonstrates best ranking ability across all risk thresholds.",
        "Calibrated Probabilities: Directly outputs mathematically sound posterior log-odds, ideal for risk-tier segmentation.",
        "High Efficiency & No Overfitting: Robust generalization without parameter bloat; serializes to < 50 KB."
    ]
    for r in reasons:
        pb = tf.add_paragraph()
        pb.text = "• " + r
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: Visual Dashboard Slide
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Performance Visualizations & Diagnostics", "Model Analytics")

    if chart_image.exists():
        slide5.shapes.add_picture(str(chart_image), Inches(0.8), Inches(1.5), width=Inches(7.2))
    
    side_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.5), Inches(4.2), Inches(5.3))
    side_box.fill.solid()
    side_box.fill.fore_color.rgb = LIGHT_BG
    side_box.line.color.rgb = RGBColor(226, 232, 240)
    tf = side_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Diagnostic Highlights\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    chart_points = [
        "F1 vs ROC-AUC: Logistic Regression balances precision and false positives better than tree models.",
        "Confusion Matrix: Successfully isolates true churners while keeping False Positives low (71.4% precision).",
        "ROC Curves: Substantial area above diagonal (AUC 0.797) proves resilient discrimination.",
        "Key Drivers: Month-to-Month contracts & low tenure are the strongest statistical drivers of churn."
    ]
    for cp in chart_points:
        pb = tf.add_paragraph()
        pb.text = "• " + cp
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 6: Business Retention Strategy & Risk Tiers
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Actionable Retention Intelligence Engine", "Business Implementation")

    # 3 Tier Cards
    t1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8))
    t1.fill.solid()
    t1.fill.fore_color.rgb = RGBColor(240, 253, 244) # Green tint
    t1.line.color.rgb = RGBColor(34, 197, 94)
    tf = t1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🟢 LOW RISK\n(< 35% Probability)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(22, 101, 52)
    bullets = [
        "Typical Profile: 2-Year Contract, High Tenure (> 24 mos), Active Tech Support.",
        "Business Action: Maintain standard cadence.",
        "Loyalty perks, referral reward programs, quarterly newsletters.",
        "Zero intervention budget needed."
    ]
    for b in bullets:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(8)

    t2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8))
    t2.fill.solid()
    t2.fill.fore_color.rgb = RGBColor(254, 252, 232) # Yellow tint
    t2.line.color.rgb = RGBColor(234, 179, 8)
    tf = t2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🟡 MODERATE RISK\n(35% - 65% Probability)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(161, 98, 7)
    bullets = [
        "Typical Profile: 1-Year contract, mid-range tenure (6-18 mos), fluctuating charges.",
        "Business Action: Automated proactive check-in.",
        "Offer free feature onboarding, tech support trial, or contract renewal discount.",
        "Low-cost preventative nudges."
    ]
    for b in bullets:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(8)

    t3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.8))
    t3.fill.solid()
    t3.fill.fore_color.rgb = RGBColor(254, 242, 242) # Red tint
    t3.line.color.rgb = RGBColor(239, 68, 68)
    tf = t3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔴 HIGH RISK\n(> 65% Probability)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(153, 27, 27)
    bullets = [
        "Typical Profile: Month-to-month, new user (< 4 mos), high charges, no support.",
        "Business Action: Immediate human intervention.",
        "Dedicated Customer Success call.",
        "Contract upgrade discount (e.g. 20% off for 1-year commitment).",
        "High ROI retention spend."
    ]
    for b in bullets:
        pb = tf.add_paragraph()
        pb.text = "• " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 7: Web Application & System Architecture
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Interactive Deployment & User Interface", "System Architecture")

    box_arch = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    box_arch.fill.solid()
    box_arch.fill.fore_color.rgb = LIGHT_BG
    box_arch.line.color.rgb = RGBColor(226, 232, 240)
    tf = box_arch.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Full Stack Deployment Architecture\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    arch_points = [
        "⚡ Lightweight Flask Microservice: Serves prediction API endpoint (`/predict_api`) with sub-50ms latency.",
        "🎨 Modern Interactive UI: Responsive frontend built with HTML5/CSS3 and asynchronous Fetch API.",
        "🔒 Zero External Network Dependency: Operates completely offline without cloud API keys, GPU instances, or subscription tokens.",
        "📦 Joblib Model Pipeline: Stores preprocessor + champion estimator in a single self-contained artifact (`best_model.joblib`).",
        "🚀 Scalability Ready: Easily wrappered into Docker, FastAPI, or AWS Lambda for high-volume enterprise production."
    ]
    for ap in arch_points:
        pb = tf.add_paragraph()
        pb.text = "• " + ap
        pb.font.size = Pt(14)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 8: Conclusion & Future Scope
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Conclusion & Future Roadmap", "Project Summary")

    c_left = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = LIGHT_BG
    c_left.line.color.rgb = RGBColor(226, 232, 240)
    tf = c_left.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Project Achievements\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    achievements = [
        "Executed full end-to-end data science lifecycle in under 2 seconds.",
        "Achieved 0.797 ROC-AUC with interpretable linear probabilistic modeling.",
        "Built multi-tier risk classification engine linked directly to business strategies.",
        "Implemented real-time web dashboard for interactive stakeholder demonstrations."
    ]
    for a in achievements:
        pb = tf.add_paragraph()
        pb.text = "✔ " + a
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(12)

    c_right = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = LIGHT_BG
    c_right.line.color.rgb = RGBColor(226, 232, 240)
    tf = c_right.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Future Enhancements\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    future = [
        "Time-Series Tracking: Monitor customer usage trajectory over weekly rolling windows.",
        "Survival Analysis: Predict precise time-to-churn (Cox Proportional Hazards model).",
        "Automated CRM Integration: Direct webhook triggers into HubSpot or Salesforce.",
        "SHAP Explanations: Add per-customer local feature attribution plots."
    ]
    for f in future:
        pb = tf.add_paragraph()
        pb.text = "➔ " + f
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_DARK
        pb.space_after = Pt(12)

    prs.save(str(output_pptx))
    print(f"[+] Successfully generated presentation deck at: {output_pptx}")

if __name__ == "__main__":
    create_deck()
